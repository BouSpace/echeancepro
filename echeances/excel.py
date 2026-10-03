"""Import et export Excel des échéances."""
import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Banque, Echeance, Fournisseur

COLONNES = [
    "Banque", "Type", "Fournisseur", "Référence", "Date création",
    "Date échéance", "Montant", "Coût aval", "Statut", "Commentaire",
]

TYPES = {
    "chèque": Echeance.Type.CHEQUE, "cheque": Echeance.Type.CHEQUE,
    "traite": Echeance.Type.TRAITE,
    # anciens libellés, acceptés pour les fichiers déjà préparés
    "traite avalisée": Echeance.Type.TRAITE, "traite avalisee": Echeance.Type.TRAITE,
    "traite simple": Echeance.Type.TRAITE,
}
S = Echeance.Statut
# Clés sans accent, en minuscules (voir _normaliser) : « Payée », « payee », « PAYÉ »… sont acceptés.
STATUTS = {
    "a venir": S.A_VENIR, "avenir": S.A_VENIR, "ouverte": S.A_VENIR, "ouvert": S.A_VENIR,
    "payee": S.PAYEE, "paye": S.PAYEE,
    "impayee": S.IMPAYEE, "impaye": S.IMPAYEE,
    "renouvelee": S.RENOUVELEE, "renouvele": S.RENOUVELEE,
    "annulee": S.ANNULEE, "annule": S.ANNULEE,
}


def _normaliser(texte):
    """Minuscules, sans accents, espaces simplifiés."""
    sans_accent = "".join(c for c in unicodedata.normalize("NFD", str(texte)) if unicodedata.category(c) != "Mn")
    return " ".join(sans_accent.lower().split())

ENTETE_FILL = PatternFill("solid", fgColor="0B4DA2")
ENTETE_FONT = Font(bold=True, color="FFFFFF")


def _entete(ws):
    ws.append(COLONNES)
    for cell in ws[1]:
        cell.fill, cell.font = ENTETE_FILL, ENTETE_FONT
        cell.alignment = Alignment(horizontal="center")
    for i, largeur in enumerate([14, 18, 32, 16, 14, 14, 16, 14, 12, 40], start=1):
        ws.column_dimensions[get_column_letter(i)].width = largeur
    ws.freeze_panes = "A2"


def modele_vierge():
    wb = Workbook()
    ws = wb.active
    ws.title = "Échéances"
    _entete(ws)
    ws.append(["CORIS", "Traite", "SAHA ENERGY", "TR-0001", date(2026, 7, 28),
               date(2026, 10, 27), 100000000, 0, "À venir", "Exemple : supprimez cette ligne"])
    _formats(ws)
    aide = wb.create_sheet("Aide")
    for ligne in [
        ["Colonne", "Valeurs acceptées"],
        ["Banque", "Code ou nom de la banque (créée automatiquement si inconnue)"],
        ["Type", "Chèque ou Traite"],
        ["Fournisseur", "Nom du fournisseur (créé automatiquement si inconnu)"],
        ["Dates", "Format date Excel ou JJ/MM/AAAA"],
        ["Montant, Coût aval", "Nombres entiers en F CFA, sans espace ni symbole"],
        ["Statut", "Facultatif : À venir (par défaut), Payée, Impayée, Renouvelée, Annulée. Les variantes sans accent ou au masculin (payé, annulé…) sont acceptées ; toute autre valeur fait refuser la ligne. Une échéance importée déjà clôturée reçoit sa date d'échéance comme date de clôture."],
    ]:
        aide.append(ligne)
    aide.column_dimensions["A"].width = 22
    aide.column_dimensions["B"].width = 80
    return wb


def _formats(ws):
    for row in ws.iter_rows(min_row=2):
        row[4].number_format = row[5].number_format = "DD/MM/YYYY"
        row[6].number_format = row[7].number_format = "#,##0"


def exporter(queryset):
    wb = Workbook()
    ws = wb.active
    ws.title = "Échéances"
    _entete(ws)
    for e in queryset.select_related("banque", "fournisseur"):
        ws.append([
            str(e.banque), e.get_type_display(), e.fournisseur.nom, e.reference,
            e.date_creation, e.date_echeance, int(e.montant), int(e.cout_aval),
            e.get_statut_display(), e.commentaire,
        ])
    _formats(ws)
    return wb


def _date(v):
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return datetime.strptime(str(v).strip(), "%d/%m/%Y").date()


def _nombre(v):
    if v in (None, ""):
        return Decimal(0)
    if isinstance(v, (int, float)):
        return Decimal(int(round(v)))
    return Decimal(str(v).replace("\u202f", "").replace("\xa0", "").replace(" ", "").replace(",", "."))


def importer(fichier, user=None):
    """Importe un fichier au format du modèle. Retourne (nb_créées, erreurs)."""
    wb = load_workbook(fichier, data_only=True)
    ws = wb.worksheets[0]
    entetes = [str(c.value).strip() if c.value else "" for c in ws[1]]
    manquantes = [c for c in COLONNES[:8] if c not in entetes]
    if manquantes:
        return 0, [f"Colonnes manquantes : {', '.join(manquantes)}. Utilisez le modèle d'import."]
    idx = {nom: entetes.index(nom) for nom in COLONNES if nom in entetes}

    a_creer, erreurs = [], []
    for num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(row):
            continue
        val = lambda col: row[idx[col]] if col in idx and idx[col] < len(row) else None  # noqa: E731
        try:
            banque_nom = str(val("Banque") or "").strip()
            fournisseur_nom = str(val("Fournisseur") or "").strip()
            type_txt = str(val("Type") or "").strip().lower()
            if not banque_nom or not fournisseur_nom:
                raise ValueError("banque et fournisseur obligatoires")
            if type_txt not in TYPES:
                raise ValueError(f"type inconnu « {val('Type')} »")
            de = _date(val("Date échéance"))
            if not de:
                raise ValueError("date d'échéance manquante")
            dc = _date(val("Date création")) or de
            montant = _nombre(val("Montant"))
            if montant <= 0:
                raise ValueError("montant nul ou négatif")
            statut_brut = val("Statut")
            if statut_brut is None or not str(statut_brut).strip():
                statut = S.A_VENIR
            else:
                statut = STATUTS.get(_normaliser(statut_brut))
                if statut is None:
                    raise ValueError(
                        f"statut inconnu « {statut_brut} » (valeurs acceptées : À venir, Payée, Impayée, Renouvelée, Annulée)")
            a_creer.append(dict(
                banque_nom=banque_nom, fournisseur_nom=fournisseur_nom, type=TYPES[type_txt],
                reference=str(val("Référence") or "").strip(), date_creation=dc, date_echeance=de,
                montant=montant, cout_aval=_nombre(val("Coût aval")), statut=statut,
                commentaire=str(val("Commentaire") or "").strip(),
            ))
            if statut != S.A_VENIR:  # échéance reprise déjà clôturée : on date la clôture à l'échéance
                a_creer[-1].update(date_cloture=de, reference_cloture="Statut repris de l'import Excel")
        except (ValueError, InvalidOperation) as exc:
            erreurs.append(f"Ligne {num} : {exc}")

    if erreurs:
        return 0, erreurs  # rien n'est importé tant qu'il reste des erreurs

    for d in a_creer:
        banque = (Banque.objects.filter(code__iexact=d["banque_nom"]).first()
                  or Banque.objects.filter(nom__iexact=d["banque_nom"]).first()
                  or Banque.objects.create(nom=d.pop("banque_nom").upper(), code=""))
        d.pop("banque_nom", None)
        fournisseur, _ = Fournisseur.objects.get_or_create(
            nom__iexact=d["fournisseur_nom"], defaults={"nom": d["fournisseur_nom"].upper()}
        )
        d.pop("fournisseur_nom")
        Echeance.objects.create(banque=banque, fournisseur=fournisseur, cree_par=user, **d)
    return len(a_creer), []
